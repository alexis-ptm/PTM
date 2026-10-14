"""
script_doc.py — a house format for script packs
======================================================
One renderer for every kind of script we hand a client or a team member:
video ad scripts, cold-call scripts, appointment-setter scripts, objection
handling, sales-call frameworks. Same document anatomy every time, so nobody
has to relearn how to read one.

Why a module and not a template file: the format is defined by code, so it
cannot drift. Change the look here and every script pack we've ever built
re-renders consistent.

The anatomy (do not reorder — see SKILL.md):

    Cover            kicker · title · subtitle · meta table
    The brief        why these scripts look the way they do (caller supplies)
    Script N         title + angle line
      meta table     name · length · whatever else identifies it
      Setup          direction: where, how, what energy
      Beats          two-column: timecode/stage │ direction + word-for-word
      Sections       labelled blocks — on-screen text, objections,
                     headline, caption, voicemail, "why this works"
    Back matter      production/filming plan, rules, how it gets judged

Usage:
    from lib.script_doc import (new_doc, cover, h1, h2, para, bullets, table,
                                callout, render_script, page_break, save)

See templates/example_pack.py for a worked example
(video ads) — a call-script pack maps onto the identical calls.
"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

# ── House palette ────────────────────────────────────────────────────────────
ACCENT = RGBColor(0x10, 0x42, 0x81)   # headings, headline lines, timecodes
INK = RGBColor(0x10, 0x1A, 0x16)      # body
MUTED = RGBColor(0x45, 0x51, 0x4B)    # direction notes, rationale
RED = RGBColor(0xB0, 0x2E, 0x2E)      # angle lines, section dividers

FILL_HEADER = "EFEEE8"
FILL_BEAT = "F6F5F0"
FILL_CALLOUT = "EEF3FA"


# ── Primitives ───────────────────────────────────────────────────────────────

def new_doc(margin=1.0):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(margin)
        s.top_margin = s.bottom_margin = Inches(0.9)
    return doc


def _shade(cell, hex_fill):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hex_fill)
    cell._tc.get_or_add_tcPr().append(el)


def para(doc, text="", size=10.5, bold=False, italic=False, color=INK,
         space_after=6, space_before=0, style=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    if text:
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.bold = bold
        r.italic = italic
        r.font.color.rgb = color
    return p


def rich(doc, chunks, size=10.5, space_after=6):
    """chunks = [(text, bold, italic, color), ...] — mixed formatting in one line."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    for text, bold, italic, color in chunks:
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.bold = bold
        r.italic = italic
        r.font.color.rgb = color
    return p


def h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    r.font.size = Pt(17)
    r.bold = True
    r.font.color.rgb = ACCENT
    return p


def h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    r.font.size = Pt(12)
    r.bold = True
    r.font.color.rgb = INK
    return p


def label(doc, text):
    """Small caps-ish section label above a block inside a script."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text.upper())
    r.font.size = Pt(8.5)
    r.bold = True
    r.font.color.rgb = MUTED
    return p


def bullets(doc, items, size=10.5):
    """items = str or (bold_lead, rest) tuples."""
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        if isinstance(it, tuple):
            head, rest = it
            r = p.add_run(head)
            r.font.size = Pt(size)
            r.bold = True
            r.font.color.rgb = INK
            r2 = p.add_run(rest)
            r2.font.size = Pt(size)
            r2.font.color.rgb = INK
        else:
            r = p.add_run(it)
            r.font.size = Pt(size)
            r.font.color.rgb = INK


def _spacer(doc):
    """Thin breather after a table.

    An empty paragraph inherits the body font, so every table used to cost a
    full blank line of vertical space. At 4pt it reads the same and buys back
    an inch on a dense page.
    """
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1
    # The paragraph mark carries its own size, and that is what sets the height
    # of an empty paragraph — shrink the mark, not just the (absent) runs.
    ppr = p._p.get_or_add_pPr()
    rpr = ppr.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        ppr.insert(0, rpr)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "8")  # half-points
    rpr.append(sz)
    return p


def _fixed(t, widths):
    """Pin the column widths.

    Cell widths alone are a suggestion — Word and Google Docs both auto-fit to
    content and ignore them, which is how a 0.7" label column ends up eating
    half the page. Fixed layout is what makes the widths real.
    """
    el = OxmlElement("w:tblLayout")
    el.set(qn("w:type"), "fixed")
    t._tbl.tblPr.append(el)
    t.autofit = False
    # tcW on every cell is what Word reads; the tblGrid is what Google Docs
    # reads on import. Set both or the columns come back evenly split.
    grid = t._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for col, w in zip(grid.findall(qn("w:gridCol")), widths):
            col.set(qn("w:w"), str(int(w * 1440)))
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Inches(w)


def _write(cell, txt, size, color=INK, space_after=2, bold_all=False):
    """Fill a cell, honouring **inline bold** and newlines.

    Bolding used to be all-or-nothing per cell, so a **phrase** in the middle of
    a sentence printed its own asterisks.
    """
    paras = str(txt).split("\n")
    for j, line in enumerate(paras):
        p = cell.paragraphs[0] if j == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        for k, chunk in enumerate(line.split("**")):
            if not chunk:
                continue
            r = p.add_run(chunk)
            r.font.size = Pt(size)
            r.bold = bold_all or k % 2 == 1
            r.font.color.rgb = color


def table(doc, headers, rows, widths=None, size=9.5, header=True):
    """Wrap any span in **double asterisks** to bold it — whole cell or inline."""
    t = doc.add_table(rows=0, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    if header:
        hdr = t.add_row().cells
        for i, htxt in enumerate(headers):
            _shade(hdr[i], FILL_HEADER)
            p = hdr[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(htxt)
            r.font.size = Pt(size)
            r.bold = True
            r.font.color.rgb = INK
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            _write(cells[i], val, size)
    if widths:
        _fixed(t, widths)
    _spacer(doc)
    return t


def callout(doc, text, color=ACCENT, fill=FILL_CALLOUT):
    """The one thing the reader must not miss. Use sparingly — one per doc."""
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    c = t.rows[0].cells[0]
    _shade(c, fill)
    c.width = Inches(6.5)
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.space_before = Pt(3)
    r = p.add_run(text)
    r.font.size = Pt(10)
    r.bold = True
    r.font.color.rgb = color
    _spacer(doc)
    return t


def page_break(doc):
    doc.add_page_break()


# ── The beat table — the heart of the format ─────────────────────────────────

def beats(doc, rows, left_width=0.8, right_width=5.7, body_size=10.5,
          note_size=9):
    """
    rows = [(marker, direction|None, spoken|None), ...]

    marker    left column — "0-3s" for video, "OPEN" / "PROBE" / "CLOSE" for a
              call, "Q1" for an interview. Short: it is a spine, not content.
    direction italic grey — what the speaker is DOING. Never words they say.
    spoken    body text — word-for-word, in quotes. Never a summary.

    The split is the whole point: a person can read only the right-hand body
    text and deliver the script cold.
    """
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for marker, direction, spoken in rows:
        cells = t.add_row().cells
        _shade(cells[0], FILL_BEAT)
        p = cells[0].paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(marker)
        r.font.size = Pt(9)
        r.bold = True
        r.font.color.rgb = ACCENT

        c = cells[1]
        first = True
        if direction:
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(direction)
            r.font.size = Pt(note_size)
            r.italic = True
            r.font.color.rgb = MUTED
            first = False
        if spoken:
            p = c.paragraphs[0] if first else c.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            for k, chunk in enumerate(str(spoken).split("**")):
                if not chunk:
                    continue
                r = p.add_run(chunk)
                r.font.size = Pt(body_size)
                r.bold = k % 2 == 1
                r.font.color.rgb = INK
    _fixed(t, [left_width, right_width])
    _spacer(doc)
    return t


# ── Cover ────────────────────────────────────────────────────────────────────

def cover(doc, kicker, title, subtitle, meta_rows):
    """meta_rows = [(label, value), ...] — prepared for / by / date / grounded in."""
    para(doc, kicker.upper(), size=9, bold=True, color=MUTED, space_after=2)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    r.font.size = Pt(26)
    r.bold = True
    r.font.color.rgb = ACCENT
    para(doc, subtitle, size=12, color=MUTED, space_after=14)
    table(doc, ["", ""], [[f"**{k}**", v] for k, v in meta_rows],
          widths=[1.5, 5.0], header=False)


# ── One script block ─────────────────────────────────────────────────────────

def render_script(doc, spec, break_after=True):
    """
    spec keys
    ---------
    n            int or str — the number shown in the heading
    title        str        — the script's name, quoted in the heading
    subtitle     str        — the ANGLE line. What job this one does, and for
                              whom. Red, bold, one line. Never decorative.
    meta         [(k, v)]   — identity table: ad name / length / CTA, or
                              call type / duration / goal / who runs it
    setup        str        — italic direction block: where, how, what energy
    beats_label  str        — defaults "Script"; use "Call flow", "Interview
                              questions", "Objection path" as appropriate
    beats        [(marker, direction, spoken)]
    sections     [(label, content, kind)] — rendered in order after the beats.
                 kind: "bullets" (list) · "body" (paragraphs, blank-line split)
                       "hero" (one bold accent line) · "muted" (rationale)
                       "table" (content = (headers, rows, widths))
    """
    h1(doc, f"Script {spec['n']} — “{spec['title']}”")
    if spec.get("subtitle"):
        para(doc, spec["subtitle"], size=9.5, bold=True, color=RED,
             space_after=8)

    if spec.get("meta"):
        table(doc, ["", ""], [[f"**{k}**", v] for k, v in spec["meta"]],
              widths=[1.2, 5.3], header=False)

    if spec.get("setup"):
        label(doc, "Setup & direction")
        para(doc, spec["setup"], size=10, italic=True, color=MUTED,
             space_after=8)

    if spec.get("beats"):
        label(doc, spec.get("beats_label", "Script"))
        beats(doc, spec["beats"])

    for lbl, content, kind in spec.get("sections", []):
        label(doc, lbl)
        if kind == "bullets":
            for line in content:
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.space_after = Pt(2)
                r = p.add_run(line)
                r.font.size = Pt(10)
                r.font.color.rgb = INK
        elif kind == "hero":
            para(doc, content, size=11, bold=True, color=ACCENT, space_after=8)
        elif kind == "muted":
            para(doc, content, size=10, color=MUTED, space_after=10)
        elif kind == "table":
            headers, rows, widths = content
            table(doc, headers, rows, widths=widths)
        else:  # "body"
            for block in content.split("\n\n"):
                para(doc, block, size=10, space_after=6)

    if break_after:
        page_break(doc)


def divider(doc, text):
    """Small red rule between groups of scripts (cold vs warm, inbound vs outbound)."""
    para(doc, text, size=9, bold=True, color=RED, space_after=2)


def save(doc, dest):
    dest = Path(dest)
    dest.parent.mkdir(parents=True, exist_ok=True)
    doc.save(dest)
    print(f"  Wrote: {dest}")
    print(f"  Size : {dest.stat().st_size / 1024:.0f} KB")
    return dest
